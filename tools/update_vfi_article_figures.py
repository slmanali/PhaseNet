#!/usr/bin/env python3
"""Заменяет в русской статье рис. 9–11 проверенными заново построенными PNG.

Требуются PNG, полученные исправленным ghosting_blur_analysis_ru_fixed.py.
Файл metadata.json проверяется, чтобы в рисунке 9 не осталось наложенной линии.
На исходные предсказания и численные метрики этот скрипт не влияет.
"""
import argparse
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from PIL import Image


def validate_figure(png, mode, index):
    png = Path(png)
    if not png.is_file():
        raise FileNotFoundError(f"Изображение не найдено: {png}")
    with Image.open(png) as im:
        w, h = im.size
        im.verify()
    if w / h <= 1.0:
        raise ValueError(f"Ожидался 4-строчный рисунок (ширина/высота > 1): {png} {w}x{h}")
    metadata_path = png.with_name(f"{mode}_{index:05d}_metadata.json")
    if not metadata_path.is_file():
        raise FileNotFoundError(
            f"Для проверки происхождения нужен {metadata_path}. "
            "Перегенерируйте изображение исправленным скриптом."
        )
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    if metadata.get("mode") != mode or metadata.get("sample_index") != index:
        raise ValueError(f"Неверный пример в {metadata_path}: ожидалось {mode} №{index}")
    if metadata.get("comparison_overlay_line") is not False:
        raise ValueError(
            f"В {metadata_path} отсутствует подтверждение, что линия НЕ нанесена на изображение. "
            "Используйте ghosting_blur_analysis_ru_fixed.py БЕЗ --show-line."
        )
    return png


def replace_picture_before_caption(doc, number, png, width_in):
    caption_index = next(
        (i for i, p in enumerate(doc.paragraphs) if p.text.strip().startswith(f"Рис. {number}.")),
        None,
    )
    if caption_index is None:
        raise ValueError(f"Подпись Рис. {number}. не найдена")
    picture_para = None
    for i in range(caption_index - 1, max(-1, caption_index - 4), -1):
        p = doc.paragraphs[i]
        if p._p.xpath('.//a:blip'):
            picture_para = p
            break
    if picture_para is None:
        raise ValueError(f"Изображение непосредственно перед Рис. {number} не найдено")
    blips = picture_para._p.xpath('.//a:blip')
    if len(blips) != 1:
        raise ValueError(f"Ожидалось одно изображение для рис. {number}, получено {len(blips)}")
    rid = blips[0].get(qn('r:embed'))
    rel = doc.part.related_parts[rid]
    rel._blob = png.read_bytes()
    with Image.open(png) as img:
        w, h = img.size
    new_w = Inches(width_in)
    new_h = int(new_w * h / w)
    for tag in ('.//wp:extent', './/a:xfrm/a:ext'):
        for extent in picture_para._p.xpath(tag):
            extent.set('cx', str(new_w))
            extent.set('cy', str(new_h))
    picture_para.paragraph_format.keep_with_next = True
    return caption_index


def replace_caption(doc, number, content):
    p = next(p for p in doc.paragraphs if p.text.strip().startswith(f"Рис. {number}."))
    if not p.runs:
        p.add_run(content)
    else:
        p.runs[0].text = content
        for run in p.runs[1:]:
            run.text = ''


def main():
    a = argparse.ArgumentParser(description=__doc__)
    a.add_argument('--article', type=Path, required=True, help='Исходная статья DOCX')
    a.add_argument('--hard', type=Path, default=Path('figures/snufilm_ghosting_blur/hard_00017_comparison.png'))
    a.add_argument('--extreme', type=Path, default=Path('figures/snufilm_ghosting_blur/extreme_00016_comparison.png'))
    a.add_argument('--profile', type=Path, default=Path('figures/snufilm_ghosting_blur/hard_00017_edge_profile.png'))
    a.add_argument('--output', type=Path, required=True, help='Новый DOCX, исходник не перезаписывается')
    args = a.parse_args()
    if args.article.resolve() == args.output.resolve():
        a.error('Для сохранности исходника --output должен отличаться от --article')
    hard = validate_figure(args.hard, 'hard', 17)
    extreme = validate_figure(args.extreme, 'extreme', 16)
    if not args.profile.is_file():
        raise FileNotFoundError(f'Русский график профиля не найден: {args.profile}')
    doc = Document(args.article)
    replace_picture_before_caption(doc, 9, hard, 6.05)
    replace_picture_before_caption(doc, 10, args.profile, 6.7)
    replace_picture_before_caption(doc, 11, extreme, 6.45)
    replace_caption(doc, 9,
        'Рис. 9. Качественное сравнение на SNU-FILM Hard, пример 17. Строки: эталон, ComplexPhaseNet, '
        'RIFE и FILM. Для всех моделей показаны полный кадр, одинаковый фрагмент движущегося пешехода '
        'и средняя абсолютная RGB-ошибка относительно эталона (фиксированная шкала 0–0,3). '
        'Для эталона ошибка равна нулю. Профили на рис. 10 вычислены вдоль общей линии '
        'с координатами (95, 475)–(205, 475); линия не нанесена на изображения.')
    replace_caption(doc, 10,
        'Рис. 10. Профили яркости вдоль общей линии на SNU-FILM Hard, пример 17. '
        'Синяя кривая — эталон, оранжевая — ComplexPhaseNet, зелёная — RIFE, красная — FILM. '
        'Координаты линии: (95, 475)–(205, 475); по горизонтали — нормированное расстояние, '
        'по вертикали — яркость в диапазоне [0, 1].')
    replace_caption(doc, 11,
        'Рис. 11. Качественное сравнение восстановления ярко освещённой вывески на SNU-FILM Extreme, '
        'пример 16. Строки: эталон, ComplexPhaseNet, RIFE и FILM. Показаны полный кадр, одинаковый '
        'увеличенный фрагмент и средняя абсолютная RGB-ошибка (фиксированная шкала 0–0,3). '
        'Сцена релевантна обсуждению изменений освещения [1,2], однако не является контролируемым '
        'экспериментом, изолирующим этот фактор.')
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    check = Document(args.output)
    for n in (9,10,11):
        assert any(p.text.startswith(f'Рис. {n}.') for p in check.paragraphs)
    print(f'Сохранена исправленная статья: {args.output}')

if __name__=='__main__':
    main()
